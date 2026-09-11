import random
import time
import os
import subprocess
import docx

def score_job(job, user_profile: str) -> int:
    score = 50
    title = job.title.lower()
    if job.job_type == "Science":
        if any(w in title for w in ["microbiome", "metagenomic", "synthetic", "scientist i"]):
            score += 30
        if "director" in title or "principal" in title: score -= 40
    else:
        if any(w in title for w in ["analyst", "associate", "venture", "investment"]):
            score += 35
        if "partner" in title or "vp" in title: score -= 30
    if "San Diego" in job.location: score += 15
    elif "Remote" in job.location: score += 10
    return min(100, max(0, score + random.randint(-5, 5)))

def set_run_text(paragraph, text):
    paragraph.text = ""
    run = paragraph.add_run(text)
    run.font.name = 'Times New Roman'

def generate_pdf_resume(job) -> str:
    time.sleep(1.5)
    
    candidates = [
        os.path.join(os.path.dirname(__file__), "..", "data", "base_resume.docx"),
        "/home/vishant/Downloads/Resumes and Cover Letters/Vishant Gandhi's Resume.docx",
        "/home/vishant/Downloads/Vishant Gandhi's Resume.docx",
        "data/base_resume.docx",
        "base_resume.docx",
    ]
    doc_path = None
    for c in candidates:
        if os.path.exists(c):
            doc_path = c
            break
            
    if not doc_path:
        raise FileNotFoundError("Base resume DOCX not found.")
        
    doc = docx.Document(doc_path)
    
    zengler_idx = -1
    pilot_idx = -1
    
    for i, p in enumerate(doc.paragraphs):
        if "Zengler Lab" in p.text:
            zengler_idx = i
        if "Licensed Private Pilot" in p.text:
            pilot_idx = i
            
    if zengler_idx != -1:
        if job.job_type == "Venture":
            set_run_text(doc.paragraphs[zengler_idx + 1], "Critically evaluated and spearheaded the strategic development of a novel commercial testing platform for urinary health, bridging the gap between bench research and market-ready therapeutics.")
            set_run_text(doc.paragraphs[zengler_idx + 2], "Collaborated with cross-functional stakeholders and managed multi-disciplinary teams to create predictive network graphs integrating 144,000+ datasets, showcasing strong analytical and project management acumen.")
        else:
            set_run_text(doc.paragraphs[zengler_idx + 1], "Leading the bioengineering of UroCom, a novel urinary tract synthetic microbial community, leveraging advanced anaerobic cell culture and omics library preparation to elucidate host-microbe interactions.")
            set_run_text(doc.paragraphs[zengler_idx + 2], "Pioneered MetaRibo-Seq and metatranscriptomics pipelines to quantify active protein translation, resulting in a co-authored publication in npj Systems Biology (2025).")

    safe_company = "".join(c if c.isalnum() else "_" for c in job.company)
    safe_title = "".join(c if c.isalnum() else "_" for c in job.title)
    
    output_dir = os.environ.get("DOWNLOADS_DIR", "")
    if not output_dir or not os.path.exists(output_dir):
        for d in ["/home/vishant/Downloads", "/tmp", os.path.join(os.path.dirname(__file__), "..", "downloads")]:
            if os.path.exists(d):
                output_dir = d
                break
        if not output_dir:
            output_dir = "/tmp"

    temp_docx = os.path.join(output_dir, f"temp_{safe_company}.docx")
    doc.save(temp_docx)
    
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        temp_docx,
        "--outdir",
        output_dir
    ]
    subprocess.run(cmd, check=True)
    
    pdf_filename = os.path.join(output_dir, f"temp_{safe_company}.pdf")
    final_pdf_filename = os.path.join(output_dir, f"Vishant_Gandhi_Resume_{safe_company}_{safe_title}.pdf")
    
    if os.path.exists(pdf_filename):
        os.rename(pdf_filename, final_pdf_filename)
    if os.path.exists(temp_docx):
        os.remove(temp_docx)
    
    return final_pdf_filename
